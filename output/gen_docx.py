from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── page setup: A4, normal margins ───────────────────────────────────────────
sec = doc.sections[0]
sec.page_width   = Cm(21.0)
sec.page_height  = Cm(29.7)
sec.top_margin    = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.left_margin   = Cm(2.54)
sec.right_margin  = Cm(2.54)

# ── styles helper ─────────────────────────────────────────────────────────────
def set_font(run, bold=False, size=11, name='Calibri'):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name

def para(text='', bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT,
         space_before=0, space_after=4, name='Calibri'):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    p.alignment = align
    if text:
        r = p.add_run(text)
        set_font(r, bold=bold, size=size, name=name)
    return p

def heading(text, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(text)
    set_font(r, bold=True, size=size)
    # underline
    r.underline = True
    return p

# ── footer with page numbers ──────────────────────────────────────────────────
def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.clear()
    # "Page X" field
    run = p.add_run('Page ')
    set_font(run, size=9)
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    r2 = OxmlElement('w:r')
    r2.append(fldChar1)
    r2.append(instrText)
    r2.append(fldChar2)
    p._p.append(r2)

add_page_number(sec)

# ═══════════════════════════════════════════════════════════════════════════════
# PAGE 1
# ═══════════════════════════════════════════════════════════════════════════════

# Header block
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(2)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('CSCU9E5 Individual Modelling Assignment (Resit)')
set_font(r, bold=True, size=14)

p2 = doc.add_paragraph()
p2.paragraph_format.space_before = Pt(0)
p2.paragraph_format.space_after  = Pt(1)
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('Computing Science and Mathematics — University of Stirling')
set_font(r2, size=10)

p3 = doc.add_paragraph()
p3.paragraph_format.space_before = Pt(0)
p3.paragraph_format.space_after  = Pt(8)
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3a = p3.add_run('Student Number: ')
set_font(r3a, size=11)
r3b = p3.add_run('YOUR_STUDENT_NUMBER')
set_font(r3b, bold=True, size=11)

# Thin rule (simulated with a bottom-bordered paragraph)
rule_p = doc.add_paragraph()
rule_p.paragraph_format.space_before = Pt(0)
rule_p.paragraph_format.space_after  = Pt(6)
pPr = rule_p._p.get_or_add_pPr()
pBdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single')
bottom.set(qn('w:sz'), '6')
bottom.set(qn('w:space'), '1')
bottom.set(qn('w:color'), '888888')
pBdr.append(bottom)
pPr.append(pBdr)

# Section heading
heading('State Machine Diagram', size=12)

# Diagram image
img_p = doc.add_paragraph()
img_p.paragraph_format.space_before = Pt(2)
img_p.paragraph_format.space_after  = Pt(4)
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_img = img_p.add_run()
run_img.add_picture('/home/user/linkedin-outreach-/output/state_machine.png',
                    width=Cm(15.5))

# Tool statement
tool_p = doc.add_paragraph()
tool_p.paragraph_format.space_before = Pt(3)
tool_p.paragraph_format.space_after  = Pt(0)
tool_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_tool = tool_p.add_run(
    'This state machine diagram was produced using Python (matplotlib library).')
set_font(r_tool, size=9)
r_tool.italic = True

# ── page break ────────────────────────────────────────────────────────────────
doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# PAGE 2
# ═══════════════════════════════════════════════════════════════════════════════

heading('State Descriptions', size=12)

states = [
    ('Active',
     'The Active state is the default running state entered immediately after the '
     'Session object is constructed. No valid authentication is currently held: this '
     'occurs either because the session has just started, because the 15-minute '
     'authentication window has expired, or because the screen was just unlocked '
     'after a lock event, requiring the user to re-authenticate before accessing '
     'the system.'),
    ('Authenticated',
     'The Authenticated state represents a session in which the user has successfully '
     'supplied the correct password within the last 15 minutes. While in this state '
     'the user may call access() repeatedly without being asked for their password '
     'again, as the guard condition time() − authTime < 15 will evaluate to true and '
     'authenticate() returns true immediately. Once 15 minutes have elapsed the '
     'session transitions back to requiring re-authentication on the next access() call.'),
    ('Locked',
     'The Locked state is a suspended state in which the screen is inaccessible. '
     'It is entered either when the user manually calls lockScreen() (which invokes '
     'lock() on the Session) from either the Active or Authenticated state, or '
     'automatically when checkPassword() returns false after a password attempt. '
     'The session object is preserved but no further access is possible until the '
     'user calls unlockScreen(), which invokes unlock() and resets authTime to −1, '
     'ensuring full re-authentication is required on the next access() call.'),
    ('Final State',
     'The Final State is the terminal pseudostate that represents the destruction of '
     'the Session object. It is reached when close() is called from any active state '
     '— Active, Authenticated, or Locked — confirming that the user may end the '
     'session at any time regardless of authentication status or screen lock. Once '
     'this state is entered the Session object is destroyed and the interaction ends.'),
]

for name, desc in states:
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(5)
    p_name.paragraph_format.space_after  = Pt(1)
    r_name = p_name.add_run(name)
    set_font(r_name, bold=True, size=11)

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(0)
    p_desc.paragraph_format.space_after  = Pt(3)
    p_desc.paragraph_format.left_indent  = Cm(0.5)
    r_desc = p_desc.add_run(desc)
    set_font(r_desc, size=11)

# small spacer
doc.add_paragraph().paragraph_format.space_after = Pt(4)

heading('Added Attributes', size=12)

attr_p = doc.add_paragraph()
attr_p.paragraph_format.space_before = Pt(4)
attr_p.paragraph_format.space_after  = Pt(2)
r_attr_name = attr_p.add_run('−authTime : Integer')
set_font(r_attr_name, bold=True, size=11)

attr_desc_p = doc.add_paragraph()
attr_desc_p.paragraph_format.space_before = Pt(0)
attr_desc_p.paragraph_format.space_after  = Pt(3)
attr_desc_p.paragraph_format.left_indent  = Cm(0.5)
attr_text = (
    'Type: Integer. '
    'Purpose: Records the value returned by time() at the moment the user last '
    'successfully authenticated. The attribute is initialised to −1 at session '
    'creation (in the T1 constructor transition) and reset to −1 whenever '
    'unlock() is called (T11), indicating that no valid authentication is currently '
    'held. It is used in the guard conditions of the authenticate() transitions: '
    'the condition time() − authTime < 15 determines whether the 15-minute window '
    'is still active (permitting re-use of the cached authentication in T6), while '
    'time() − authTime ≥ 15 determines that the window has elapsed and full '
    're-authentication is required (T7 and T8).'
)
r_attr_desc = attr_desc_p.add_run(attr_text)
set_font(r_attr_desc, size=11)

out = '/home/user/linkedin-outreach-/output/CSCU9E5_StateMachine.docx'
doc.save(out)
print(f'Saved: {out}')
